#!/usr/bin/env python3
"""Create editable Word alternatives for every Delirium Downloads sheet."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from generate_pdfs import ROOT, Sheet, ascii_hyphens, parse_master


OUTPUT_DIR = ROOT / "output" / "word"
PUBLIC_DIR = ROOT / "public" / "downloads"

NAVY = "102A43"
NAVY_SOFT = "334E68"
BLUE = "4F77FF"
BLUE_PALE = "DFE7FF"
CORAL = "FF6B5D"
CORAL_PALE = "FFD9D4"
LIME = "C9F45B"
PAPER = "FFFDF8"
GREY_LINE = "B7C3CC"

# compact_reference_guide with named UK_A4_LOCAL_TEMPLATE overrides:
# A4 portrait; 17.5 mm margins; Arial family; branded title/callout colours.
CONTENT_WIDTH_DXA = 9920


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "244FC5")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([colour, underline])
    text = OxmlElement("w:t")
    text.text = label
    run.append(run_properties)
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_markdown(paragraph, text: str) -> None:
    text = ascii_hyphens(text)
    token_pattern = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*)")
    cursor = 0
    for match in token_pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run)
        token = match.group(0)
        link_match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
        if link_match:
            add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, italic=True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_styles(document: Document, kind: str) -> None:
    body_size = 11 if kind == "family" else 10.75
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(body_size)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6 if kind == "family" else 5)
    normal.paragraph_format.line_spacing = 1.25 if kind == "family" else 1.22

    for style_name, size, colour, before, after, font_name in (
        ("Heading 1", 16, NAVY, 18, 10, "Georgia"),
        ("Heading 2", 13, NAVY, 14, 7, "Georgia"),
        ("Heading 3", 12, NAVY_SOFT, 10, 5, "Arial"),
    ):
        style = document.styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:ascii"), font_name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    bullet.font.size = Pt(body_size)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4 if kind == "family" else 3.5)
    bullet.paragraph_format.line_spacing = 1.25 if kind == "family" else 1.22


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.text = "DELIRIUM DOWNLOADS  ·  EDITABLE LOCAL-ADAPTATION TEMPLATE"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.runs[0]
    set_run_font(run, size=8, color=NAVY_SOFT, bold=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), WD_TAB_ALIGNMENT.RIGHT)
    left_run = footer_paragraph.add_run(
        "deliriumdownloads.com · Template v0.9 · 7 August 2026 · Human clinical review required"
    )
    set_run_font(left_run, size=7.5, color=NAVY_SOFT)
    footer_paragraph.add_run("\t")
    run = footer_paragraph.add_run("Page ")
    set_run_font(run, size=7.5, color=NAVY_SOFT)
    add_page_number(footer_paragraph)


def add_title_block(document: Document, sheet: Sheet) -> None:
    banner = document.add_paragraph()
    banner.paragraph_format.space_before = Pt(4)
    banner.paragraph_format.space_after = Pt(8)
    banner_run = banner.add_run("EDITABLE TEMPLATE  ·  CLINICAL REVIEW REQUIRED BEFORE USE")
    set_run_font(banner_run, size=8.5, color=NAVY, bold=True)
    p_pr = banner._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CORAL_PALE if sheet.kind == "family" else BLUE_PALE)
    p_pr.append(shading)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run(sheet.title)
    set_run_font(title_run, name="Georgia", size=23, color=NAVY, bold=True)

    audience = document.add_paragraph()
    audience.paragraph_format.space_after = Pt(10)
    label = "PATIENT & FAMILY HANDOUT" if sheet.kind == "family" else sheet.audience.upper()
    audience_run = audience.add_run(f"INTENDED USE: {label}")
    set_run_font(audience_run, size=8.5, color=BLUE, bold=True)

    purpose = document.add_paragraph()
    purpose.paragraph_format.space_after = Pt(12)
    purpose_run = purpose.add_run(sheet.purpose)
    set_run_font(purpose_run, size=11.5, color=NAVY_SOFT, italic=True)


def add_local_adaptation_fields(document: Document) -> None:
    table = document.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, [2500, 7420], indent_dxa=120)
    header = table.cell(0, 0).merge(table.cell(0, 1))
    header.text = "LOCAL ADAPTATION FIELDS - replace highlighted text before use"
    set_cell_shading(header, LIME)
    set_repeat_table_header(table.rows[0])
    header.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in header.paragraphs[0].runs:
        set_run_font(run, size=9, color=NAVY, bold=True)

    fields = [
        ("Organisation / service", "[ADD LOCAL NAME OR LOGO]"),
        ("Clinical lead / owner", "[ADD NAME, ROLE AND CONTACT]"),
        ("Escalation or referral route", "[ADD LOCAL ROUTE AND TELEPHONE / BLEEP DETAILS]"),
        ("Local policy / review note", "[ADD POLICY LINK, FORMULARY NOTE OR NEXT REVIEW DATE]"),
    ]
    for row_index, (label, value) in enumerate(fields, start=1):
        label_cell, value_cell = table.rows[row_index].cells
        label_cell.text = label
        value_cell.text = value
        label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(label_cell, "EEF2F5")
        set_cell_shading(value_cell, PAPER)
        set_run_font(label_cell.paragraphs[0].runs[0], size=9, color=NAVY, bold=True)
        set_run_font(value_cell.paragraphs[0].runs[0], size=9, color=NAVY_SOFT, italic=True)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_safety_box(document: Document, text: str, family: bool) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(6)
    paragraph.paragraph_format.right_indent = Pt(6)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CORAL_PALE if family else BLUE_PALE)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "10")
        border.set(qn("w:space"), "5")
        border.set(qn("w:color"), NAVY)
        borders.append(border)
    p_pr.append(borders)
    add_inline_markdown(paragraph, text)
    for run in paragraph.runs:
        set_run_font(run, size=10.5, color=NAVY)


def add_content(document: Document, sheet: Sheet) -> None:
    for section in sheet.sections:
        heading = document.add_paragraph(section.heading, style="Heading 2")
        heading.paragraph_format.keep_with_next = True
        is_sources = section.heading.lower().startswith("verified")
        for block_type, data in section.blocks:
            if block_type == "paragraph":
                paragraph = document.add_paragraph()
                add_inline_markdown(paragraph, str(data))
            elif block_type == "bullets":
                for item in data:
                    paragraph = document.add_paragraph(style="List Bullet")
                    add_inline_markdown(paragraph, item)
                    if is_sources:
                        paragraph.paragraph_format.space_after = Pt(2)
                        for run in paragraph.runs:
                            set_run_font(run, size=9.5)
            elif block_type == "safety":
                add_safety_box(document, str(data), sheet.kind == "family")


def add_disclosure(document: Document) -> None:
    heading = document.add_paragraph("Publication and local review", style="Heading 2")
    heading.paragraph_format.keep_with_next = True
    paragraph = document.add_paragraph()
    text = (
        "UK-focused general education. This editable file is supplied so a health or care organisation can add "
        "local contacts and pathways. Local edits may change clinical meaning and must be reviewed by an appropriately "
        "qualified named clinician before use. It does not replace individual assessment, current national guidance, "
        "local policy or formulary advice. Prepared from source materials curated by Professor Alasdair MacLullich, "
        "University of Edinburgh; AI-assisted production. Professor MacLullich led development of the 4AT and reports "
        "no financial interest in its uptake. Final named human clinical approval is required before use in patient "
        "care. Independent resource; not an official University of Edinburgh or NHS publication."
    )
    add_inline_markdown(paragraph, text)


def set_document_language(document: Document) -> None:
    styles = document.styles.element
    doc_defaults = styles.find(qn("w:docDefaults"))
    if doc_defaults is None:
        doc_defaults = OxmlElement("w:docDefaults")
        styles.insert(0, doc_defaults)
    r_pr_default = doc_defaults.find(qn("w:rPrDefault"))
    if r_pr_default is None:
        r_pr_default = OxmlElement("w:rPrDefault")
        doc_defaults.append(r_pr_default)
    r_pr = r_pr_default.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        r_pr_default.append(r_pr)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-GB")


def build_document(sheet: Sheet, output_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.75)
    section.bottom_margin = Cm(1.75)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    configure_styles(document, sheet.kind)
    set_document_language(document)
    add_header_footer(document)
    add_title_block(document, sheet)
    add_local_adaptation_fields(document)
    add_content(document, sheet)
    add_disclosure(document)

    properties = document.core_properties
    properties.title = f"{sheet.title} - editable local-adaptation template"
    properties.subject = "Delirium education resource - human clinical review required"
    properties.author = "Professor Alasdair MacLullich"
    properties.keywords = "delirium, local adaptation, clinical education"
    properties.comments = "AI-assisted production. Final named human clinical review required."

    document.save(output_path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    PUBLIC_DIR.mkdir(parents=True, exist_ok=True)
    sheets = parse_master(ROOT / "content" / "family-sheets.md", "family")
    sheets += parse_master(ROOT / "content" / "professional-sheets.md", "professional")
    sheets += parse_master(ROOT / "content" / "additional-sheets.md", "additional")
    if len(sheets) != 25:
        raise SystemExit(f"Expected 25 sheets, found {len(sheets)}")
    for sheet in sheets:
        output_path = OUTPUT_DIR / f"{sheet.slug}.docx"
        build_document(sheet, output_path)
        shutil.copy2(output_path, PUBLIC_DIR / output_path.name)
    print(f"Generated {len(sheets)} editable Word templates")


if __name__ == "__main__":
    main()
